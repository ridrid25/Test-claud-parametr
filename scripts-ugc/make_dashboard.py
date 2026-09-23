# Продукт №2: «Прибыль селлера за 15 минут» — каскад прибыли из детализации
# еженедельного отчёта WB. Правила проекта: сводная сходится до рубля с
# контрольной суммой; компенсации в деньгах остаются, в штуках не считаются;
# незнакомые формулировки падают в «Прочее (проверить)»; все строки открыты,
# нули видны, в шапке нейтральная воронка автофильтра с подписью.
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEAL = "0A655E"; PURPLE = "A01A6E"; RED = "C00000"; GREEN = "1E7B1E"
YELLOW = PatternFill("solid", fgColor="FFF2A8")
DARK = PatternFill("solid", fgColor="0A655E")
LIGHT = PatternFill("solid", fgColor="F2F7F6")
RUB = '#,##0.00" ₽";[Red]-#,##0.00" ₽";0.00" ₽"'
PCT = '0.0%;[Red]-0.0%;0.0%'
QTY = '#,##0;[Red]-#,##0;0'

def f(size=12, bold=False, color="000000"):
    return Font(name="Times New Roman", size=size, bold=bold, color=color)

wb = Workbook()

# ---------------- Лист 1: Данные ----------------
wd = wb.active
wd.title = "Данные"
wd.sheet_view.showGridLines = False
widths = {"A": 44, "B": 12, "C": 18, "D": 22, "E": 16, "F": 14, "G": 16}
for col, w in widths.items():
    wd.column_dimensions[col].width = w

wd.merge_cells("A1:G1")
c = wd["A1"]; c.value = "ШАГ 1. Вставьте сюда строки из детализации еженедельного отчёта WB"
c.font = f(14, True, "FFFFFF"); c.fill = DARK
c.alignment = Alignment(horizontal="center", vertical="center")
wd.row_dimensions[1].height = 26
wd.merge_cells("A2:G2")
c = wd["A2"]
c.value = ("Нужны 7 колонок из отчёта (значения ниже — пример, замените своими). "
           "Строки добавляйте сколько угодно — каскад считает диапазон до 5000 строк.")
c.font = f(12); c.alignment = Alignment(wrap_text=True, vertical="center")
wd.row_dimensions[2].height = 30

headers = ["Обоснование для оплаты", "Кол-во, шт", "Продано (Пр), ₽",
           "К перечислению за товар, ₽", "Логистика, ₽", "Штрафы, ₽", "Удержания, ₽"]
for i, h in enumerate(headers):
    c = wd.cell(row=3, column=i + 1, value=h)
    c.font = f(12, True); c.fill = LIGHT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
wd.row_dimensions[3].height = 30
wd.auto_filter.ref = "A3:G3"

example = [
    ["Продажа", 118, 175840, 135397, 0, 0, 0],
    ["Продажа", 2, 0, 2400, 0, 0, 0],
    ["Возврат", 6, 8940, 6883, 0, 0, 0],
    ["Логистика", 0, 0, 0, 9840, 0, 0],
    ["Штраф", 0, 0, 0, 0, 1500, 0],
    ["Удержание", 0, 0, 0, 0, 0, 2100],
    ["Добровольная компенсация при возврате", 1, 0, 800, 0, 0, 0],
    ["Оказание услуг «ВБ.Продвижение»", 0, 0, 0, 0, 0, 4200],
]
thin = Side(style="thin", color="BFBFBF")
for r0, row in enumerate(example, start=4):
    for cix, v in enumerate(row, start=1):
        c = wd.cell(row=r0, column=cix, value=v)
        c.font = f(12); c.fill = YELLOW
        c.number_format = QTY if cix == 2 else (RUB if cix > 2 else "General")
        c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        c.alignment = Alignment(horizontal="left" if cix == 1 else "right", vertical="center")

wd.cell(row=13, column=1,
        value="Примечание: строка «Продажа» с нулевым «Продано» — компенсационная: "
              "ВБ оформляет её как продажу с ценой 0. В штуках каскад её не считает, в деньгах — оставляет.")
wd["A13"].font = f(12); wd["A13"].alignment = Alignment(wrap_text=True)
wd.merge_cells("A13:G13"); wd.row_dimensions[13].height = 30

D = "Данные"
RR = "4:5000"  # рабочий диапазон
A = f"{D}!$A$4:$A$5000"; B = f"{D}!$B$4:$B$5000"; C_ = f"{D}!$C$4:$C$5000"
Dc = f"{D}!$D$4:$D$5000"; E = f"{D}!$E$4:$E$5000"; F_ = f"{D}!$F$4:$F$5000"; G = f"{D}!$G$4:$G$5000"

# ---------------- Лист 2: Каскад прибыли ----------------
ws = wb.create_sheet("Каскад прибыли")
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 52
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 46

def cell(addr, value, font=None, fill=None, fmt=None, align=None, wrap=False):
    c = ws[addr]; c.value = value; c.font = font or f()
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    c.alignment = Alignment(horizontal=align or ("left" if isinstance(value, str) else "right"),
                            vertical="center", wrap_text=wrap)
    return c

ws.merge_cells("B2:D2")
cell("B2", "ПРИБЫЛЬ СЕЛЛЕРА ЗА 15 МИНУТ — КАСКАД ПО ОТЧЁТУ WB",
     f(16, True, "FFFFFF"), DARK, align="center")
ws.row_dimensions[2].height = 30
ws.merge_cells("B3:D3")
cell("B3", "Все статьи всегда на месте, нулевые не удаляются — в новом периоде статья с суммой "
           "появится сама. Скрыть нулевые строки можно воронкой автофильтра ▼ в шапке таблицы.",
     f(12), wrap=True)
ws.row_dimensions[3].height = 32

# Шапка каскада с автофильтром
cell("B5", "Статья", f(12, True), LIGHT, align="left")
cell("C5", "Сумма, ₽", f(12, True), LIGHT, align="right")
cell("D5", "Как посчитано", f(12, True), LIGHT, align="left")
ws.auto_filter.ref = "B5:D5"

r = 6
cell(f"B{r}", "Продано товара (Пр)")
cell(f"C{r}", f'=SUMIFS({C_},{A},"Продажа")-SUMIFS({C_},{A},"Возврат")', fmt=RUB)
cell(f"D{r}", "продажи минус возвраты, в ценах реализации", f(12)); sold = f"C{r}"; r += 1
cell(f"B{r}", "Вознаграждение ВВ (комиссия)", f(12, False, RED))
cell(f"C{r}", f'={sold}-(SUMIFS({Dc},{A},"Продажа")-SUMIFS({Dc},{A},"Возврат"))', f(12, False, RED), fmt=RUB)
cell(f"D{r}", "разница между Пр и «к перечислению за товар»", f(12)); r += 1
cell(f"B{r}", "К перечислению за товар", f(12, True))
cell(f"C{r}", f'=SUMIFS({Dc},{A},"Продажа")-SUMIFS({Dc},{A},"Возврат")', f(12, True), fmt=RUB)
tovar = f"C{r}"; r += 1
cell(f"B{r}", "Компенсации", f(12, False, GREEN))
cell(f"C{r}", f'=SUMPRODUCT(ISNUMBER(SEARCH("компенсац",{A}))*{Dc})', f(12, False, GREEN), fmt=RUB)
cell(f"D{r}", "в деньгах остаются; в штуках не считаются", f(12)); comp = f"C{r}"; r += 1
cell(f"B{r}", "Логистика", f(12, False, RED))
cell(f"C{r}", f"=-SUM({E})", f(12, False, RED), fmt=RUB); logi = f"C{r}"; r += 1
cell(f"B{r}", "Штрафы", f(12, False, RED))
cell(f"C{r}", f"=-SUM({F_})", f(12, False, RED), fmt=RUB); fine = f"C{r}"; r += 1
cell(f"B{r}", "Удержания", f(12, False, RED))
cell(f"C{r}", f"=-SUM({G})", f(12, False, RED), fmt=RUB); hold = f"C{r}"; r += 1
cell(f"B{r}", "Прочее (проверить!)")
cell(f"C{r}", f'=SUMPRODUCT(({A}<>"Продажа")*({A}<>"Возврат")*(1-ISNUMBER(SEARCH("компенсац",{A})))*({A}<>"")*{Dc})', fmt=RUB)
cell(f"D{r}", "незнакомые формулировки падают сюда, а не ломают расчёт", f(12), wrap=True)
other = f"C{r}"; r += 1
cell(f"B{r}", "К ПЕРЕЧИСЛЕНИЮ (РАСЧЁТ)", f(14, True, TEAL), LIGHT)
cell(f"C{r}", f"={tovar}+{comp}+{logi}+{fine}+{hold}+{other}", f(14, True, TEAL), LIGHT, fmt=RUB)
calc = f"C{r}"; ws.row_dimensions[r].height = 24; r += 1
cell(f"B{r}", "К перечислению по отчёту WB (впишите итог из отчёта)")
cell(f"C{r}", 114074, fill=YELLOW, fmt=RUB); rep = f"C{r}"; r += 1
cell(f"B{r}", "РАСХОЖДЕНИЕ (должно быть 0,00 ₽)", f(12, True))
cell(f"C{r}", f"={calc}-{rep}", f(12, True, RED), fmt=RUB)
cell(f"D{r}", "расхождение — это ошибка, а не «примерно сошлось»", f(12)); r += 2

cell(f"B{r}", "ЭКОНОМИКА ПЕРИОДА", f(13, True, PURPLE)); r += 1
cell(f"B{r}", "Продано, шт (компенсационные строки не считаются)")
cell(f"C{r}", f'=SUMIFS({B},{A},"Продажа",{C_},">0")-SUMIFS({B},{A},"Возврат",{C_},">0")', fmt=QTY)
qty = f"C{r}"; r += 1
cell(f"B{r}", "Себестоимость единицы (закупка + доставка до вас)")
cell(f"C{r}", 420, fill=YELLOW, fmt=RUB); unit = f"C{r}"; r += 1
cell(f"B{r}", "Себестоимость проданного", f(12, False, RED))
cell(f"C{r}", f"=-{qty}*{unit}", f(12, False, RED), fmt=RUB); cogs = f"C{r}"; r += 1
cell(f"B{r}", "Ставка налога (УСН «доходы» — от Пр)")
cell(f"C{r}", 0.06, fill=YELLOW, fmt=PCT); taxr = f"C{r}"; r += 1
cell(f"B{r}", "Налог", f(12, False, RED))
cell(f"C{r}", f"=-{sold}*{taxr}", f(12, False, RED), fmt=RUB); tax = f"C{r}"; r += 1
cell(f"B{r}", "ПРИБЫЛЬ ПЕРИОДА", f(14, True, TEAL), LIGHT)
cell(f"C{r}", f"={calc}+{cogs}+{tax}", f(14, True, TEAL), LIGHT, fmt=RUB)
profit = f"C{r}"; ws.row_dimensions[r].height = 24; r += 1
cell(f"B{r}", "Маржинальность (прибыль / Пр)", f(12, True))
cell(f"C{r}", f"=IF({sold}=0,0,{profit}/{sold})", f(12, True, TEAL), fmt=PCT); r += 2

ws.merge_cells(f"B{r}:D{r}")
cell(f"B{r}", "Автоматически, по всем неделям и карточкам сразу — на mp.ridfinance.ru. "
              "Вопросы: @RidFinancebot_bot", f(12, True), wrap=True)
ws.row_dimensions[r].height = 28

for row in ws.iter_rows(min_row=5, max_row=r - 2, min_col=3, max_col=3):
    for c in row:
        if c.value is not None:
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)

for sheet in (wd, ws):
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

wb.save("Дашборд_Прибыль_селлера_WB.xlsx")
print("saved")
